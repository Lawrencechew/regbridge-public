import hashlib
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.outputs.errors import OutputGenerationError
from app.outputs.models import OutputIssueCode
from app.outputs.service import safe_filename
from app.outputs.xlsx import XLSX_MEDIA_TYPE, XlsxTemplateGenerator
from tests.output_helpers import output_dataset, output_definition, write_output_template


def generate(tmp_path, *, name="Alpha"):
    template = tmp_path / "template.xlsx"
    digest = write_output_template(template)
    definition = output_definition(template, digest)
    content = XlsxTemplateGenerator().generate(
        output_dataset(name=name), definition, template, 1024 * 1024
    )
    return template, content


def test_scalar_table_types_and_template_structure_are_preserved(tmp_path) -> None:
    template, content = generate(tmp_path)
    original = template.read_bytes()
    workbook = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    assert workbook["Summary"]["C1"].value == "Alpha"
    assert workbook["Summary"]["A1"].value == "unchanged"
    assert workbook["Summary"]["A1"].font.bold is True
    assert workbook["Summary"]["B1"].value == "=1+1"
    assert str(workbook["Summary"].merged_cells) == "D1:E1"
    assert workbook["Hidden"].sheet_state == "hidden"
    row = workbook["Table"][2]
    assert row[0].value == "Alpha" and row[0].data_type == "s"
    assert row[1].value == 7 and row[1].data_type == "n"
    assert row[2].value == 12.5 and row[2].data_type == "n"
    assert row[3].value is True and row[3].data_type == "b"
    assert type(row[4].value) is datetime and row[4].value.date() == date(2026, 8, 14)
    assert type(row[5].value) is datetime
    workbook.close()
    assert template.read_bytes() == original


@pytest.mark.parametrize(
    "dangerous",
    ['=1+1', '+1+1', '-1+1', '@SUM(A1:A2)', '=HYPERLINK("http://example.com")'],
)
def test_formula_injection_values_reload_as_text(tmp_path, dangerous) -> None:
    _, content = generate(tmp_path, name=dangerous)
    workbook = load_workbook(BytesIO(content), data_only=False)
    for cell in (workbook["Summary"]["C1"], workbook["Table"]["A2"]):
        assert cell.value == dangerous
        assert cell.data_type == "s"
    workbook.close()


def test_null_writes_a_blank_cell(tmp_path) -> None:
    template = tmp_path / "template.xlsx"
    digest = write_output_template(template)
    definition = output_definition(
        template,
        digest,
        mappings=[
            {
                "type": "scalar_cell",
                "source": {"type": "literal", "value": None},
                "target": {"sheet": "Summary", "cell": "C1"},
            }
        ],
    )
    content = XlsxTemplateGenerator().generate(output_dataset(), definition, template, 1024 * 1024)
    workbook = load_workbook(BytesIO(content))
    assert workbook["Summary"]["C1"].value is None


def test_runtime_template_hash_is_rechecked(tmp_path) -> None:
    template = tmp_path / "template.xlsx"
    digest = write_output_template(template)
    definition = output_definition(template, digest)
    template.write_bytes(template.read_bytes() + b"modified")
    with pytest.raises(OutputGenerationError) as caught:
        XlsxTemplateGenerator().generate(output_dataset(), definition, template, 1024 * 1024)
    assert caught.value.code == OutputIssueCode.OUTPUT_TEMPLATE_INTEGRITY_FAILED


def test_generated_size_limit_is_enforced(tmp_path) -> None:
    template = tmp_path / "template.xlsx"
    digest = write_output_template(template)
    with pytest.raises(OutputGenerationError) as caught:
        XlsxTemplateGenerator().generate(output_dataset(), output_definition(template, digest), template, 10)
    assert caught.value.code == OutputIssueCode.OUTPUT_FILE_TOO_LARGE


@pytest.mark.parametrize("name", ["../evil.xlsx", "..\\evil.xlsx", "C:\\evil.xlsx", "/evil.xlsx"])
def test_filename_rejects_escaping_paths(name) -> None:
    with pytest.raises(OutputGenerationError):
        safe_filename(name, pack_id="pack", pack_version="1.0.0", dataset_id="data")


def test_filename_is_bounded_and_xlsx_only() -> None:
    with pytest.raises(OutputGenerationError):
        safe_filename("x" * 181 + ".xlsx", pack_id="p", pack_version="1.0.0", dataset_id="d")
    with pytest.raises(OutputGenerationError):
        safe_filename("output.csv", pack_id="p", pack_version="1.0.0", dataset_id="d")
    assert safe_filename("{pack_id}-{dataset_id}.xlsx", pack_id="p", pack_version="1.0.0", dataset_id="d") == "p-d.xlsx"


def test_generated_bytes_are_valid_and_hashable(tmp_path) -> None:
    _, content = generate(tmp_path)
    assert len(content) > 0
    assert len(hashlib.sha256(content).hexdigest()) == 64
    assert XLSX_MEDIA_TYPE.endswith("spreadsheetml.sheet")


def test_record_table_preserves_dataset_order_across_multiple_rows(tmp_path) -> None:
    template = tmp_path / "template.xlsx"
    digest = write_output_template(template)
    definition = output_definition(template, digest)
    dataset = output_dataset()
    second = dataset.records[0].model_copy(
        update={
            "record_id": "record-2",
            "values": {
                **dataset.records[0].values,
                "name": dataset.records[0].values["name"].model_copy(update={"value": "Beta"}),
            },
        }
    )
    dataset = dataset.model_copy(update={"records": [dataset.records[0], second]})
    # The unique scalar is intentionally removed for this table-only ordering check.
    table_only = definition.model_copy(update={"mappings": (definition.mappings[1],)})
    content = XlsxTemplateGenerator().generate(dataset, table_only, template, 1024 * 1024)
    workbook = load_workbook(BytesIO(content))
    assert [workbook["Table"][f"A{row}"].value for row in (2, 3)] == ["Alpha", "Beta"]
    workbook.close()
