import hashlib
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from app.canonical.models import (
    CanonicalDataType,
    CanonicalDataset,
    CanonicalFieldDefinition,
    CanonicalRecord,
    CanonicalSchema,
    CanonicalValue,
)
from app.outputs.models import OutputDefinition


def output_schema() -> CanonicalSchema:
    fields = (
        ("name", CanonicalDataType.STRING),
        ("count", CanonicalDataType.INTEGER),
        ("amount", CanonicalDataType.DECIMAL),
        ("expected_amount", CanonicalDataType.DECIMAL),
        ("active", CanonicalDataType.BOOLEAN),
        ("event_date", CanonicalDataType.DATE),
        ("processed_at", CanonicalDataType.DATETIME),
        ("group_code", CanonicalDataType.STRING),
        ("row_key", CanonicalDataType.STRING),
    )
    return CanonicalSchema(
        id="synthetic-output",
        version="1.0.0",
        name="Synthetic Output",
        fields=tuple(
            CanonicalFieldDefinition(id=field_id, name=field_id, data_type=data_type)
            for field_id, data_type in fields
        ),
    )


def output_dataset(*, name: str = "Alpha", amount: Decimal = Decimal("12.50")) -> CanonicalDataset:
    schema = output_schema()
    values = {
        "name": name,
        "count": 7,
        "amount": amount,
        "expected_amount": amount,
        "active": True,
        "event_date": date(2026, 8, 14),
        "processed_at": datetime(2026, 8, 14, 9, 30),
        "group_code": "A",
        "row_key": "First",
    }
    record = CanonicalRecord(
        record_id="record-1",
        values={
            field_id: CanonicalValue(
                data_type=schema.fields_by_id[field_id].data_type,
                value=value,
            )
            for field_id, value in values.items()
        },
    )
    return CanonicalDataset(dataset_id="synthetic-dataset", schema=schema, records=[record])


def write_output_template(path: Path) -> str:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = "unchanged"
    sheet["A1"].font = Font(bold=True)
    sheet["B1"] = "=1+1"
    sheet.merge_cells("D1:E1")
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    table = workbook.create_sheet("Table")
    table.append(["Name", "Count", "Amount", "Active", "Date", "Datetime", "Key"])
    table.append([None] * 6 + ["First"])
    workbook.save(path)
    workbook.close()
    return hashlib.sha256(path.read_bytes()).hexdigest()


def output_definition(path: Path, digest: str, *, mappings=None) -> OutputDefinition:
    return OutputDefinition.model_validate(
        {
            "id": "synthetic-xlsx",
            "version": "1.0.0",
            "type": "xlsx_template",
            "template": {
                "source_id": "SRC-001",
                "filename": path.name,
                "sha256": digest,
                "local_path": path.name,
            },
            "filename": {"pattern": "output-{dataset_id}.xlsx"},
            "mappings": mappings
            or [
                {
                    "type": "scalar_cell",
                    "source": {"type": "dataset_field", "field_id": "name"},
                    "target": {"sheet": "Summary", "cell": "C1"},
                },
                {
                    "type": "record_table",
                    "records": {"where": [{"field_id": "group_code", "operator": "equals", "value": "A"}]},
                    "sheet": "Table",
                    "start_row": 2,
                    "end_row": 10,
                    "columns": [
                        {"field_id": "name", "column": "A"},
                        {"field_id": "count", "column": "B"},
                        {"field_id": "amount", "column": "C"},
                        {"field_id": "active", "column": "D"},
                        {"field_id": "event_date", "column": "E"},
                        {"field_id": "processed_at", "column": "F"},
                    ],
                },
            ],
        },
        strict=False,
    )
