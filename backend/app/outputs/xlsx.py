from __future__ import annotations

import hashlib
import zipfile
import warnings
from copy import copy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path, PurePosixPath, PureWindowsPath

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.workbook import Workbook

from app.canonical.models import CanonicalDataset, CanonicalRecord, CanonicalSchema
from app.outputs.errors import OutputGenerationError
from app.outputs.models import (
    DatasetFieldSource,
    MatrixCellMapping,
    OutputConditionOperator,
    OutputDefinition,
    OutputIssueCode,
    OutputScalar,
    RecordSelection,
    RecordTableMapping,
    RuntimeFieldSource,
    ScalarCellMapping,
)
from app.regpacks.errors import RegPackComponentError

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def resolve_template_path(pack_root: Path, local_path: str) -> Path:
    posix = PurePosixPath(local_path)
    windows = PureWindowsPath(local_path)
    if posix.is_absolute() or windows.is_absolute() or ".." in posix.parts or ".." in windows.parts:
        raise RegPackComponentError("The output template path is unsafe.")
    candidate = pack_root / local_path
    if candidate.is_symlink():
        raise RegPackComponentError("Output template symlinks are not allowed.")
    resolved = candidate.resolve()
    if not resolved.is_relative_to(pack_root.resolve()):
        raise RegPackComponentError("The output template path escapes the RegPack.")
    return resolved


def validate_xlsx_definition(
    definition: OutputDefinition,
    schema: CanonicalSchema | dict[str, CanonicalSchema],
    pack_root: Path,
    source_ids: set[str],
    runtime_field_ids: set[str] | None = None,
) -> Path:
    if definition.template.source_id not in source_ids:
        raise RegPackComponentError("The output template has a dangling source reference.")
    if not definition.template.filename.lower().endswith(".xlsx"):
        raise RegPackComponentError("Only macro-free XLSX output templates are supported.")
    template_path = resolve_template_path(pack_root, definition.template.local_path)
    if not template_path.is_file():
        raise RegPackComponentError("The output template is missing.")
    if template_path.stat().st_size > 20 * 1024 * 1024:
        raise RegPackComponentError("The output template exceeds the safe size limit.")
    if template_path.name != definition.template.filename:
        raise RegPackComponentError("The output template filename metadata does not match.")
    digest = hashlib.sha256(template_path.read_bytes()).hexdigest()
    if digest != definition.template.sha256:
        raise RegPackComponentError("The output template failed its integrity check.")
    try:
        with zipfile.ZipFile(template_path) as archive:
            names = {name.casefold() for name in archive.namelist()}
            if any(name.endswith("vbaproject.bin") for name in names):
                raise RegPackComponentError("Macro-enabled output templates are not supported.")
    except zipfile.BadZipFile as exc:
        raise RegPackComponentError("The output template is not a valid XLSX archive.") from exc
    schemas = schema if isinstance(schema, dict) else {"default": schema}
    runtime_field_ids = runtime_field_ids or set()
    for mapping in definition.mappings:
        section_id = getattr(mapping, "section_id", None)
        if isinstance(mapping, ScalarCellMapping) and isinstance(mapping.source, DatasetFieldSource):
            section_id = mapping.source.section_id
        if isinstance(mapping, ScalarCellMapping) and isinstance(mapping.source, RuntimeFieldSource):
            if mapping.source.field_id not in runtime_field_ids:
                raise RegPackComponentError("The output definition references an unknown runtime field.")
            continue
        if section_id is None:
            if len(schemas) != 1:
                raise RegPackComponentError("Sectioned output mappings require a section_id.")
            selected_schema = next(iter(schemas.values()))
        else:
            selected_schema = schemas.get(section_id)
            if selected_schema is None:
                raise RegPackComponentError("The output definition references an unknown section.")
        referenced = set()
        if isinstance(mapping, ScalarCellMapping) and isinstance(mapping.source, DatasetFieldSource):
            referenced.add(mapping.source.field_id)
            if mapping.source.selection:
                referenced.update(item.field_id for item in mapping.source.selection.where)
        elif isinstance(mapping, RecordTableMapping):
            referenced.update(item.field_id for item in mapping.records.where)
            referenced.update(column.field_id for column in mapping.columns if column.field_id)
            if mapping.row_lookup:
                referenced.add(mapping.row_lookup.field_id)
        elif isinstance(mapping, MatrixCellMapping):
            referenced.update(item.field_id for item in mapping.records.where)
            referenced.update(
                (mapping.value_field_id, mapping.row_lookup.field_id, mapping.column_lookup.field_id)
            )
        if not referenced.issubset(set(selected_schema.fields_by_id)):
            raise RegPackComponentError("The output definition references an unknown canonical field.")

    try:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
            workbook = load_workbook(template_path, data_only=False, keep_links=False, read_only=False)
    except Exception as exc:
        raise RegPackComponentError("The output template is not a valid XLSX workbook.") from exc
    try:
        if getattr(workbook, "_external_links", []):
            raise RegPackComponentError("Output templates with external links are not supported.")
        claimed: set[tuple[str, str]] = set()
        for mapping in definition.mappings:
            sheet_name = mapping.target.sheet if isinstance(mapping, ScalarCellMapping) else mapping.sheet
            if sheet_name not in workbook.sheetnames:
                raise RegPackComponentError("The output definition references an unknown worksheet.")
            sheet = workbook[sheet_name]
            for coordinate in _declared_coordinates(mapping):
                key = (sheet_name, coordinate)
                if key in claimed:
                    raise RegPackComponentError("The output definition contains conflicting targets.")
                claimed.add(key)
                _assert_writable(sheet, coordinate, component_error=True)
            if isinstance(mapping, RecordTableMapping) and mapping.row_lookup:
                _build_row_lookup(sheet, mapping.row_lookup, component_error=True)
            if isinstance(mapping, MatrixCellMapping):
                _build_row_lookup(sheet, mapping.row_lookup, component_error=True)
                _build_column_lookup(sheet, mapping.column_lookup, component_error=True)
    finally:
        workbook.close()
    return template_path


def _declared_coordinates(mapping) -> list[str]:
    if isinstance(mapping, ScalarCellMapping):
        return [mapping.target.cell]
    if isinstance(mapping, RecordTableMapping):
        if mapping.row_lookup:
            rows = range(mapping.row_lookup.start_row, mapping.row_lookup.end_row + 1)
            excluded = set(mapping.row_lookup.excluded_rows)
        else:
            end = mapping.end_row if mapping.end_row is not None else mapping.start_row
            rows = range(mapping.start_row, end + 1)
            excluded = set()
        return [f"{column.column}{row}" for row in rows if row not in excluded for column in mapping.columns]
    rows = range(mapping.row_lookup.start_row, mapping.row_lookup.end_row + 1)
    excluded = set(mapping.row_lookup.excluded_rows)
    start = column_index_from_string(mapping.column_lookup.start_column)
    end = column_index_from_string(mapping.column_lookup.end_column)
    if end < start:
        raise RegPackComponentError("A matrix column lookup range is reversed.")
    return [f"{get_column_letter(column)}{row}" for row in rows if row not in excluded for column in range(start, end + 1)]


def _assert_writable(sheet, coordinate: str, *, component_error: bool = False) -> None:
    cell = sheet[coordinate]
    invalid = isinstance(cell, MergedCell) or cell.data_type == "f" or (
        sheet.protection.sheet and cell.protection.locked
    )
    if invalid:
        if component_error:
            raise RegPackComponentError("The output definition targets a non-writable cell.")
        raise OutputGenerationError(
            OutputIssueCode.OUTPUT_TARGET_NOT_WRITABLE,
            "An output mapping targets a protected, merged, or calculated cell.",
        )


def _build_row_lookup(sheet, lookup, *, component_error: bool = False) -> dict[object, int]:
    result: dict[object, int] = {}
    excluded = set(lookup.excluded_rows)
    for row in range(lookup.start_row, lookup.end_row + 1):
        if row in excluded:
            continue
        value = sheet[f"{lookup.template_column}{row}"].value
        if value is None:
            continue
        if value in result:
            error = "The template row lookup contains duplicate keys."
            if component_error:
                raise RegPackComponentError(error)
            raise OutputGenerationError(OutputIssueCode.OUTPUT_MAPPING_INVALID, error)
        result[value] = row
    return result


def _build_column_lookup(sheet, lookup, *, component_error: bool = False) -> dict[object, int]:
    start = column_index_from_string(lookup.start_column)
    end = column_index_from_string(lookup.end_column)
    if end < start:
        if component_error:
            raise RegPackComponentError("A matrix column lookup range is reversed.")
        raise OutputGenerationError(OutputIssueCode.OUTPUT_MAPPING_INVALID, "A matrix column lookup is invalid.")
    result: dict[object, int] = {}
    for column in range(start, end + 1):
        value = sheet.cell(lookup.template_row, column).value
        if value is None:
            continue
        if value in result:
            if component_error:
                raise RegPackComponentError("The template column lookup contains duplicate keys.")
            raise OutputGenerationError(OutputIssueCode.OUTPUT_MAPPING_INVALID, "A template column lookup is ambiguous.")
        result[value] = column
    return result


class XlsxTemplateGenerator:
    def generate(
        self,
        dataset: CanonicalDataset,
        definition: OutputDefinition,
        template_path: Path,
        max_file_size_bytes: int,
    ) -> bytes:
        return self.generate_bundle(
            {"default": dataset}, definition, template_path, max_file_size_bytes, {}
        )

    def generate_bundle(
        self,
        datasets: dict[str, CanonicalDataset],
        definition: OutputDefinition,
        template_path: Path,
        max_file_size_bytes: int,
        runtime_values: dict[str, object],
    ) -> bytes:
        self._verify_template(template_path, definition)
        try:
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
                workbook = load_workbook(template_path, data_only=False, keep_links=False, read_only=False)
        except Exception as exc:
            raise OutputGenerationError(OutputIssueCode.OUTPUT_TEMPLATE_INVALID, "The verified output template could not be loaded.") from exc
        written: set[tuple[str, str]] = set()
        try:
            for mapping in definition.mappings:
                if isinstance(mapping, ScalarCellMapping):
                    if isinstance(mapping.source, RuntimeFieldSource):
                        value = runtime_values.get(mapping.source.field_id)
                    else:
                        dataset = self._dataset(datasets, getattr(mapping.source, "section_id", None))
                        value = self._scalar_value(dataset, mapping.source)
                    self._write(workbook[mapping.target.sheet], mapping.target.cell, value, written)
                elif isinstance(mapping, RecordTableMapping):
                    dataset = self._dataset(datasets, mapping.section_id)
                    self._write_table(workbook, dataset, mapping, written)
                else:
                    dataset = self._dataset(datasets, mapping.section_id)
                    self._write_matrix(workbook, dataset, mapping, written)
            output = BytesIO()
            workbook.save(output)
            content = _restore_worksheet_extensions(template_path.read_bytes(), output.getvalue())
        except OutputGenerationError:
            raise
        except Exception as exc:
            raise OutputGenerationError(OutputIssueCode.OUTPUT_GENERATION_FAILED, "The output workbook could not be generated.") from exc
        finally:
            workbook.close()
        if len(content) > max_file_size_bytes:
            raise OutputGenerationError(OutputIssueCode.OUTPUT_FILE_TOO_LARGE, "The generated artifact exceeds the configured size limit.")
        try:
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
                check = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
                check.close()
        except Exception as exc:
            raise OutputGenerationError(OutputIssueCode.OUTPUT_GENERATION_FAILED, "The generated XLSX failed structural verification.") from exc
        return content

    @staticmethod
    def _dataset(
        datasets: dict[str, CanonicalDataset], section_id: str | None
    ) -> CanonicalDataset:
        if section_id is not None:
            dataset = datasets.get(section_id)
            if dataset is None:
                raise OutputGenerationError(
                    OutputIssueCode.OUTPUT_MAPPING_INVALID,
                    "An output mapping references a missing canonical section.",
                )
            return dataset
        if len(datasets) != 1:
            raise OutputGenerationError(
                OutputIssueCode.OUTPUT_MAPPING_INVALID,
                "A sectioned output mapping is missing section identity.",
            )
        return next(iter(datasets.values()))

    @staticmethod
    def _verify_template(path: Path, definition: OutputDefinition) -> None:
        if not path.is_file():
            raise OutputGenerationError(OutputIssueCode.OUTPUT_TEMPLATE_NOT_FOUND, "The configured output template was not found.")
        if hashlib.sha256(path.read_bytes()).hexdigest() != definition.template.sha256:
            raise OutputGenerationError(OutputIssueCode.OUTPUT_TEMPLATE_INTEGRITY_FAILED, "The output template failed its SHA-256 integrity check.")

    def _write_table(self, workbook: Workbook, dataset: CanonicalDataset, mapping: RecordTableMapping, written: set[tuple[str, str]]) -> None:
        sheet = workbook[mapping.sheet]
        records = [record for record in dataset.records if _matches(record, mapping.records)]
        lookup = _build_row_lookup(sheet, mapping.row_lookup) if mapping.row_lookup else None
        for offset, record in enumerate(records):
            if lookup is not None:
                key = _record_value(record, mapping.row_lookup.field_id)
                if key not in lookup:
                    raise OutputGenerationError(OutputIssueCode.OUTPUT_TARGET_NOT_FOUND, "A canonical record does not match a declared template row.")
                row = lookup[key]
            else:
                row = mapping.start_row + offset
                if mapping.end_row is not None and row > mapping.end_row:
                    raise OutputGenerationError(OutputIssueCode.OUTPUT_CAPACITY_EXCEEDED, "The output table does not have enough declared rows.")
                if mapping.copy_row_style and offset and mapping.start_row:
                    self._copy_row_style(sheet, mapping.start_row, row, mapping.columns)
            for column in mapping.columns:
                if column.sequence:
                    value = offset + 1
                else:
                    value = column.literal if column.field_id is None else _record_value(record, column.field_id)
                self._write(sheet, f"{column.column}{row}", value, written)

    def _write_matrix(self, workbook: Workbook, dataset: CanonicalDataset, mapping: MatrixCellMapping, written: set[tuple[str, str]]) -> None:
        sheet = workbook[mapping.sheet]
        rows = _build_row_lookup(sheet, mapping.row_lookup)
        columns = _build_column_lookup(sheet, mapping.column_lookup)
        for record in dataset.records:
            if not _matches(record, mapping.records):
                continue
            row_key = _record_value(record, mapping.row_lookup.field_id)
            column_key = _record_value(record, mapping.column_lookup.field_id)
            if row_key not in rows or column_key not in columns:
                raise OutputGenerationError(OutputIssueCode.OUTPUT_TARGET_NOT_FOUND, "A canonical record does not match a declared template matrix target.")
            coordinate = f"{get_column_letter(columns[column_key])}{rows[row_key]}"
            self._write(sheet, coordinate, _record_value(record, mapping.value_field_id), written)

    @staticmethod
    def _copy_row_style(sheet, source_row: int, target_row: int, columns) -> None:
        for column in columns:
            source = sheet[f"{column.column}{source_row}"]
            target = sheet[f"{column.column}{target_row}"]
            if source.has_style:
                target._style = copy(source._style)
            target.number_format = source.number_format
            target.protection = copy(source.protection)
            target.alignment = copy(source.alignment)

    @staticmethod
    def _write(sheet, coordinate: str, value: OutputScalar, written: set[tuple[str, str]]) -> None:
        key = (sheet.title, coordinate)
        if key in written:
            raise OutputGenerationError(OutputIssueCode.OUTPUT_MAPPING_INVALID, "Multiple records or mappings resolved to the same output cell.")
        _assert_writable(sheet, coordinate)
        cell = sheet[coordinate]
        try:
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"
        except Exception as exc:
            raise OutputGenerationError(OutputIssueCode.OUTPUT_VALUE_SERIALISATION_FAILED, "A canonical value could not be serialised to XLSX.") from exc
        written.add(key)

    @staticmethod
    def _scalar_value(dataset: CanonicalDataset, source) -> OutputScalar:
        if source.type == "literal":
            return source.value
        records = dataset.records
        if source.selection:
            records = [record for record in records if _matches(record, source.selection)]
        values = [_record_value(record, source.field_id) for record in records]
        distinct = []
        for value in values:
            if value not in distinct:
                distinct.append(value)
        if len(distinct) != 1:
            raise OutputGenerationError(OutputIssueCode.OUTPUT_MAPPING_INVALID, "A dataset_field scalar source must resolve to exactly one unique value.")
        return distinct[0]


def _record_value(record: CanonicalRecord, field_id: str) -> OutputScalar:
    value = record.values.get(field_id)
    return None if value is None else value.value


def _matches(record: CanonicalRecord, selection: RecordSelection) -> bool:
    for condition in selection.where:
        actual = _record_value(record, condition.field_id)
        if condition.operator == OutputConditionOperator.EQUALS and actual != condition.value:
            return False
        if condition.operator == OutputConditionOperator.NOT_EQUALS and actual == condition.value:
            return False
        if condition.operator == OutputConditionOperator.IS_NULL and actual is not None:
            return False
        if condition.operator == OutputConditionOperator.IS_NOT_NULL and actual is None:
            return False
    return True


def _restore_worksheet_extensions(template: bytes, generated: bytes) -> bytes:
    """Restore unsupported worksheet extension lists removed by openpyxl."""
    from xml.etree import ElementTree

    main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    replacements: dict[str, bytes] = {}
    with zipfile.ZipFile(BytesIO(template)) as source:
        with zipfile.ZipFile(BytesIO(generated)) as current:
            for name in current.namelist():
                if not name.startswith("xl/worksheets/") or not name.endswith(".xml"):
                    continue
                try:
                    source_root = ElementTree.fromstring(source.read(name))
                except KeyError:
                    continue
                extensions = source_root.find(f"{main}extLst")
                if extensions is None:
                    continue
                target_root = ElementTree.fromstring(current.read(name))
                existing = target_root.find(f"{main}extLst")
                if existing is not None:
                    target_root.remove(existing)
                target_root.append(extensions)
                replacements[name] = ElementTree.tostring(
                    target_root, encoding="utf-8", xml_declaration=False
                )
            output = BytesIO()
            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                for info in current.infolist():
                    archive.writestr(info, replacements.get(info.filename, current.read(info.filename)))
    return output.getvalue()
