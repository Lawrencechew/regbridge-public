import csv
import hashlib
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import PurePosixPath, PureWindowsPath
from xml.etree.ElementTree import iterparse
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from app.imports.errors import ImportRequestError
from app.imports.models import (
    ImportIssueCode,
    SourceColumn,
    SourceFileInspection,
    SourcePreviewRow,
    SourceSheet,
    SourceType,
    json_safe_source_value,
)
from app.imports.security import (
    ImportLimits,
    enforce_file_size,
    inspect_xlsx_container,
    source_type_for_name,
)


@dataclass(frozen=True)
class SourceCell:
    value: object | None
    is_formula: bool = False


@dataclass(frozen=True)
class SourceDataRow:
    source_row: int
    cells: tuple[SourceCell, ...]


@dataclass(frozen=True)
class SourceTable:
    inspection: SourceFileInspection
    rows: tuple[SourceDataRow, ...]
    merged_ranges: tuple[tuple[int, int, int, int], ...] = ()


class SourceInspector:
    def __init__(self, limits: ImportLimits) -> None:
        self.limits = limits

    def inspect(
        self,
        source_name: str,
        content: bytes,
        *,
        sheet_name: str | None = None,
        header_row: int = 1,
    ) -> SourceTable:
        if not 1 <= header_row <= 50:
            raise ImportRequestError(
                ImportIssueCode.INVALID_HEADER_ROW,
                "Header row must be between 1 and 50.",
                422,
            )
        safe_name = PureWindowsPath(PurePosixPath(source_name).name).name or "upload"
        source_type = source_type_for_name(safe_name)
        enforce_file_size(content, self.limits)
        file_sha256 = hashlib.sha256(content).hexdigest()
        if source_type == SourceType.CSV:
            if sheet_name is not None:
                raise ImportRequestError(
                    ImportIssueCode.SHEET_NOT_FOUND,
                    "CSV sources do not contain worksheets.",
                    422,
                )
            return self._inspect_csv(safe_name, content, file_sha256, header_row)
        return self._inspect_xlsx(
            safe_name, content, file_sha256, sheet_name, header_row
        )

    def inspect_many(
        self,
        source_name: str,
        content: bytes,
        requests: tuple[tuple[str, int], ...],
    ) -> tuple[SourceTable, ...]:
        """Inspect several XLSX sheets while parsing the uploaded package once."""
        if not requests:
            return ()
        if any(not 1 <= header_row <= 50 for _, header_row in requests):
            raise ImportRequestError(
                ImportIssueCode.INVALID_HEADER_ROW,
                "Header row must be between 1 and 50.",
                422,
            )
        safe_name = PureWindowsPath(PurePosixPath(source_name).name).name or "upload"
        source_type = source_type_for_name(safe_name)
        enforce_file_size(content, self.limits)
        if source_type != SourceType.XLSX:
            if len(requests) != 1:
                raise ImportRequestError(
                    ImportIssueCode.SHEET_NOT_FOUND,
                    "Multi-section imports require an XLSX workbook.",
                    422,
                )
            _, header_row = requests[0]
            return (self.inspect(safe_name, content, header_row=header_row),)
        file_sha256 = hashlib.sha256(content).hexdigest()
        has_external_links = inspect_xlsx_container(content, self.limits)
        try:
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
        except Exception as exc:
            raise ImportRequestError(
                ImportIssueCode.INVALID_XLSX,
                "The XLSX workbook could not be parsed safely.",
                400,
            ) from exc
        try:
            if len(workbook.worksheets) > self.limits.max_xlsx_sheets:
                raise ImportRequestError(
                    ImportIssueCode.SOURCE_LIMIT_EXCEEDED,
                    "The workbook exceeds the configured worksheet limit.",
                    413,
                )
            sheets = tuple(
                SourceSheet(
                    name=sheet.title,
                    visibility=sheet.sheet_state,
                    row_count=max(sheet.max_row - min(row for _, row in requests), 0),
                    column_count=sheet.max_column,
                )
                for sheet in workbook.worksheets
            )
            for sheet in sheets:
                self._enforce_table_limits(sheet.row_count, sheet.column_count)
            results = tuple(
                self._inspect_loaded_sheet(
                    workbook,
                    safe_name,
                    content,
                    file_sha256,
                    sheets,
                    sheet_name,
                    header_row,
                    has_external_links,
                )
                for sheet_name, header_row in requests
            )
            return results
        finally:
            workbook.close()

    def _inspect_loaded_sheet(
        self,
        workbook,
        source_name: str,
        content: bytes,
        file_sha256: str,
        sheets: tuple[SourceSheet, ...],
        sheet_name: str,
        header_row: int,
        has_external_links: bool,
    ) -> SourceTable:
        if sheet_name not in workbook.sheetnames:
            raise ImportRequestError(
                ImportIssueCode.SHEET_NOT_FOUND,
                "A requested workflow worksheet was not found.",
                422,
            )
        worksheet = workbook[sheet_name]
        if worksheet.sheet_state != "visible":
            raise ImportRequestError(
                ImportIssueCode.SHEET_NOT_FOUND,
                "Hidden worksheets cannot be selected as customer data sections.",
                422,
            )
        if header_row > worksheet.max_row:
            raise ImportRequestError(
                ImportIssueCode.INVALID_HEADER_ROW,
                "The requested header row is outside the source table.",
                422,
            )
        raw_rows = worksheet.iter_rows(
            min_row=header_row,
            max_row=worksheet.max_row,
            max_col=worksheet.max_column,
        )
        header_cells = next(raw_rows)
        headers = tuple(
            None if cell.value is None or cell.data_type == "f" else str(cell.value)
            for cell in header_cells
        )
        source_rows = tuple(
            SourceDataRow(
                source_row=row_number,
                cells=tuple(
                    SourceCell(value=cell.value, is_formula=cell.data_type == "f")
                    for cell in row
                ),
            )
            for row_number, row in enumerate(raw_rows, start=header_row + 1)
        )
        columns = tuple(
            SourceColumn(index=index, header=header, excel_column=get_column_letter(index))
            for index, header in enumerate(headers, start=1)
        )
        inspection = SourceFileInspection(
            source_name=source_name,
            source_type=SourceType.XLSX,
            file_size_bytes=len(content),
            file_sha256=file_sha256,
            sheets=sheets,
            selected_sheet=sheet_name,
            selection_required=False,
            header_row=header_row,
            columns=columns,
            data_row_count=len(source_rows),
            preview_rows=self._preview(source_rows),
            formula_cells_detected=(
                any(cell.data_type == "f" for cell in header_cells)
                or any(cell.is_formula for row in source_rows for cell in row.cells)
            ),
            warnings=(
                ("External workbook links are present and were not followed.",)
                if has_external_links
                else ()
            ),
        )
        return SourceTable(
            inspection=inspection,
            rows=source_rows,
            merged_ranges=self._merged_ranges(content, worksheet._worksheet_path),
        )

    def _inspect_csv(
        self, source_name: str, content: bytes, file_sha256: str, header_row: int
    ) -> SourceTable:
        if b"\x00" in content:
            raise ImportRequestError(
                ImportIssueCode.UNSUPPORTED_TEXT_ENCODING,
                "The CSV source contains binary content.",
                400,
            )
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ImportRequestError(
                ImportIssueCode.UNSUPPORTED_TEXT_ENCODING,
                "The CSV source must use UTF-8 encoding.",
                400,
            ) from exc
        try:
            rows = list(csv.reader(StringIO(text), delimiter=",", strict=True))
        except csv.Error as exc:
            raise ImportRequestError(
                ImportIssueCode.UNSUPPORTED_TEXT_ENCODING,
                "The CSV source could not be parsed safely.",
                400,
            ) from exc
        if header_row > len(rows):
            raise ImportRequestError(
                ImportIssueCode.INVALID_HEADER_ROW,
                "The requested header row is outside the source table.",
                422,
            )
        column_count = max((len(row) for row in rows), default=0)
        data_count = max(len(rows) - header_row, 0)
        self._enforce_table_limits(data_count, column_count)
        headers = self._headers(rows[header_row - 1], column_count)
        source_rows = tuple(
            SourceDataRow(
                source_row=index,
                cells=tuple(
                    SourceCell(row[column] if column < len(row) else "")
                    for column in range(column_count)
                ),
            )
            for index, row in enumerate(rows[header_row:], start=header_row + 1)
        )
        columns = tuple(
            SourceColumn(index=index, header=header, excel_column=None)
            for index, header in enumerate(headers, start=1)
        )
        inspection = SourceFileInspection(
            source_name=source_name,
            source_type=SourceType.CSV,
            file_size_bytes=len(content),
            file_sha256=file_sha256,
            sheets=(),
            selected_sheet=None,
            selection_required=False,
            header_row=header_row,
            columns=columns,
            data_row_count=data_count,
            preview_rows=self._preview(source_rows),
            formula_cells_detected=False,
        )
        return SourceTable(inspection=inspection, rows=source_rows)

    def _inspect_xlsx(
        self,
        source_name: str,
        content: bytes,
        file_sha256: str,
        requested_sheet: str | None,
        header_row: int,
    ) -> SourceTable:
        has_external_links = inspect_xlsx_container(content, self.limits)
        try:
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
        except Exception as exc:
            raise ImportRequestError(
                ImportIssueCode.INVALID_XLSX,
                "The XLSX workbook could not be parsed safely.",
                400,
            ) from exc
        try:
            if len(workbook.worksheets) > self.limits.max_xlsx_sheets:
                raise ImportRequestError(
                    ImportIssueCode.SOURCE_LIMIT_EXCEEDED,
                    "The workbook exceeds the configured worksheet limit.",
                    413,
                )
            sheets = tuple(
                SourceSheet(
                    name=sheet.title,
                    visibility=sheet.sheet_state,
                    row_count=max(sheet.max_row - header_row, 0),
                    column_count=sheet.max_column,
                )
                for sheet in workbook.worksheets
            )
            for sheet in sheets:
                self._enforce_table_limits(sheet.row_count, sheet.column_count)

            visible_names = [
                sheet.name for sheet in sheets if sheet.visibility == "visible"
            ]
            selected = requested_sheet
            if selected is not None and selected not in workbook.sheetnames:
                raise ImportRequestError(
                    ImportIssueCode.SHEET_NOT_FOUND,
                    "The requested worksheet was not found.",
                    422,
                )
            if selected is None and len(visible_names) == 1:
                selected = visible_names[0]
            selection_required = selected is None and len(visible_names) != 1
            if selection_required:
                warnings = ["Select a visible worksheet to inspect its columns."]
                if has_external_links:
                    warnings.append(
                        "External workbook links are present and were not followed."
                    )
                inspection = SourceFileInspection(
                    source_name=source_name,
                    source_type=SourceType.XLSX,
                    file_size_bytes=len(content),
                    file_sha256=file_sha256,
                    sheets=sheets,
                    selected_sheet=None,
                    selection_required=True,
                    header_row=header_row,
                    columns=(),
                    data_row_count=None,
                    preview_rows=(),
                    formula_cells_detected=self._contains_formula(
                        content,
                        tuple(
                            sheet._worksheet_path for sheet in workbook.worksheets
                        ),
                    ),
                    warnings=tuple(warnings),
                )
                return SourceTable(inspection=inspection, rows=())

            worksheet = workbook[selected]
            if header_row > worksheet.max_row:
                raise ImportRequestError(
                    ImportIssueCode.INVALID_HEADER_ROW,
                    "The requested header row is outside the source table.",
                    422,
                )
            raw_rows = worksheet.iter_rows(
                min_row=header_row,
                max_row=worksheet.max_row,
                max_col=worksheet.max_column,
            )
            header_cells = next(raw_rows)
            header_formula_detected = any(
                cell.data_type == "f" for cell in header_cells
            )
            headers = tuple(
                None
                if cell.value is None or cell.data_type == "f"
                else str(cell.value)
                for cell in header_cells
            )
            source_rows = tuple(
                SourceDataRow(
                    source_row=row_number,
                    cells=tuple(
                        SourceCell(value=cell.value, is_formula=cell.data_type == "f")
                        for cell in row
                    ),
                )
                for row_number, row in enumerate(raw_rows, start=header_row + 1)
            )
            merged_ranges = self._merged_ranges(content, worksheet._worksheet_path)
            columns = tuple(
                SourceColumn(
                    index=index,
                    header=header,
                    excel_column=get_column_letter(index),
                )
                for index, header in enumerate(headers, start=1)
            )
            inspection = SourceFileInspection(
                source_name=source_name,
                source_type=SourceType.XLSX,
                file_size_bytes=len(content),
                file_sha256=file_sha256,
                sheets=sheets,
                selected_sheet=selected,
                selection_required=False,
                header_row=header_row,
                columns=columns,
                data_row_count=len(source_rows),
                preview_rows=self._preview(source_rows),
                formula_cells_detected=header_formula_detected
                or any(cell.is_formula for row in source_rows for cell in row.cells),
                warnings=(
                    ("External workbook links are present and were not followed.",)
                    if has_external_links
                    else ()
                ),
            )
            return SourceTable(
                inspection=inspection,
                rows=source_rows,
                merged_ranges=merged_ranges,
            )
        finally:
            workbook.close()

    def _enforce_table_limits(self, rows: int, columns: int) -> None:
        if rows > self.limits.max_rows or columns > self.limits.max_columns:
            raise ImportRequestError(
                ImportIssueCode.SOURCE_LIMIT_EXCEEDED,
                "The source table exceeds a configured row or column limit.",
                413,
            )

    @staticmethod
    def _headers(row: list[str], column_count: int) -> tuple[str | None, ...]:
        return tuple(
            (row[index] if index < len(row) and row[index] != "" else None)
            for index in range(column_count)
        )

    def _preview(
        self, rows: tuple[SourceDataRow, ...]
    ) -> tuple[SourcePreviewRow, ...]:
        return tuple(
            SourcePreviewRow(
                source_row=row.source_row,
                values=tuple(
                    {"formula": True}
                    if cell.is_formula
                    else json_safe_source_value(cell.value)
                    for cell in row.cells
                ),
            )
            for row in rows[: self.limits.preview_rows]
        )

    @staticmethod
    def _merged_ranges(
        content: bytes, worksheet_path: str
    ) -> tuple[tuple[int, int, int, int], ...]:
        ranges: list[tuple[int, int, int, int]] = []
        try:
            with ZipFile(BytesIO(content)) as archive:
                with archive.open(worksheet_path) as sheet_xml:
                    for _, element in iterparse(sheet_xml, events=("start",)):
                        if element.tag.endswith("mergeCell"):
                            reference = element.attrib.get("ref")
                            if reference:
                                ranges.append(range_boundaries(reference))
        except (KeyError, ValueError):
            return ()
        return tuple(ranges)

    @staticmethod
    def _contains_formula(content: bytes, worksheet_paths: tuple[str, ...]) -> bool:
        try:
            with ZipFile(BytesIO(content)) as archive:
                for worksheet_path in worksheet_paths:
                    with archive.open(worksheet_path) as sheet_xml:
                        for _, element in iterparse(sheet_xml, events=("start",)):
                            if element.tag.endswith("}f") or element.tag == "f":
                                return True
        except (KeyError, ValueError):
            return False
        return False
